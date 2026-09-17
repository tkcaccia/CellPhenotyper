#!/usr/bin/env python3
"""Add the final UNI2-grid and boundary-refinement development results to the DOCX pair."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.dml.color import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


REVISION_DATE = datetime(2026, 9, 15, tzinfo=timezone.utc)


def find_exact(document: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph equal to {text!r}, found {len(matches)}")
    return matches[0]


def find_starting(document: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def replace_starting(document: Document, prefix: str, text: str) -> Paragraph:
    paragraph = find_starting(document, prefix)
    paragraph.text = text
    return paragraph


def replace_exact(document: Document, original: str, text: str) -> Paragraph:
    paragraph = find_exact(document, original)
    paragraph.text = text
    return paragraph


def insert_before(anchor: Paragraph, text: str, style: str) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def set_revision_metadata(document: Document, *, subject: str, description: str) -> None:
    props = document.core_properties
    props.subject = subject
    props.description = description
    props.modified = REVISION_DATE
    try:
        props.revision = max(2, int(props.revision or "0") + 1)
    except ValueError:
        props.revision = 2


def normalize_heading_presentation(document: Document) -> None:
    """Keep the retained typography while enforcing plain black document headings."""
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.color.rgb = RGBColor(0, 0, 0)
            if style_name == "Title":
                properties = style.element.get_or_add_pPr()
                borders = properties.find(qn("w:pBdr"))
                if borders is not None:
                    properties.remove(borders)
    for paragraph in document.paragraphs:
        if paragraph.style.name not in {"Title", "Heading 1", "Heading 2", "Heading 3"}:
            continue
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
        if paragraph.style.name == "Title":
            properties = paragraph._p.get_or_add_pPr()
            borders = properties.find(qn("w:pBdr"))
            if borders is not None:
                properties.remove(borders)


def revise_main(source: Path, destination: Path) -> None:
    document = Document(source)

    results_anchor = find_exact(
        document,
        "The complete run is feasible on a single 16-GB GPU workstation",
    )
    insert_before(results_anchor, "Development selection of spatial resolution and boundary regularization", "Heading 2")
    insert_before(
        results_anchor,
        "A separate development experiment evaluated spatial-grid UNI-2 sampling and final vector regularization on a 43,458 x 39,886-pixel breast-cancer tissue image with a three-class expert annotation. The annotation was withheld from pipeline inference and was used only after candidate generation. With K fixed at three for this comparison, a 28-pixel central token-pooling square and a 56-pixel grid stride produced a mean class intersection over union (IoU) of 0.9454 and mean Dice score of 0.9717 after label matching. The corresponding values were 0.8837 and 0.9383 for a 56-pixel square with the same stride, and 0.6817 and 0.8077 for the historical coupled 90-pixel geometry. Both 56-pixel-stride routes generated 498,992 observations. Summed task runtime was 35,247.2 s for the 28/56 route and 34,925.6 s for the 56/56 route, a difference of 321.6 s (0.92%). These measurements support the 28/56 geometry for grid-based tissue-domain development, but they do not establish performance on independent specimens.",
        "Normal",
    )
    insert_before(
        results_anchor,
        "The selected final vector procedure applied one selective shared-boundary Laplacian pass after topology-preserving simplification. It reduced the 95th percentile boundary-turn angle from 121.13 degrees to 116.58 degrees and the density of long sharp turns from 1,355.0 to 1,292.4 per 100,000 ring pixels, without changing the external tissue footprint, feature count or 244,005-segment count. The maximum vertex displacement was 11.856 native pixels, and positive-area overlap between classes remained zero. Mean class IoU changed from 0.94536 to 0.94234 and boundary F1 at 32 pixels from 0.91979 to 0.91938. We therefore treat this as a development-selected geometric regularizer that improves angular plausibility with a small loss of agreement to this reference, rather than as evidence of improved segmentation accuracy.",
        "Normal",
    )

    replace_starting(
        document,
        "Before MedSAM inference, adjacent KODAMA tissue labels compete",
        "Before MedSAM inference, adjacent KODAMA tissue labels compete within a bounded internal-boundary error band. The optimizer uses deterministic mean-field annealing of a multiclass Potts objective comprising robustly scaled Lab and optical-density appearance terms and edge-weighted spatial smoothness. The default eight-neighbour system includes distance-weighted diagonal interactions to reduce grid-aligned staircasing. At each cooling iteration, a pixel can retain its current label or receive only a label present on its local frontier, which prevents non-local label transfer. Temperature decreases geometrically from 2.0 to 0.05 over 16 iterations at fourfold working downsampling and a 64-native-pixel boundary radius. The lowest hard-energy state observed is retained. Confident eroded label cores, foreground/background membership and GrandQC clean-tissue support remain invariant. MedSAM receives this optimized categorical baseline together with the original protected seeds. The operation is unsupervised and does not use the expert development annotation.",
    )

    replace_starting(
        document,
        "The final multiclass GeoJSON is derived",
        "The final multiclass GeoJSON is derived from the native-resolution MedSAM-refined categorical raster using four-connected rasterio polygon topology. Interior rings and every positive-area annotated component are retained. Shared boundaries with unequal collinear vertices are noded once and polygonized into atomic faces that retain their original raster class. Shared-boundary Visvalingam-Whyatt simplification is optimized up to a configured maximum tolerance subject to an exact categorical symmetric-difference limit against the unsimplified coverage; a zero limit requests exact geometry. Per-class hole filling and buffering remain disabled. After simplification, one endpoint-preserving Laplacian pass selectively moves internal vertices only when the turn is at least 45 degrees and both adjacent segments are at least 16 native pixels. The displacement coefficient is 0.05. The external tissue footprint is fixed, no vertices are added, and the modified linework is polygonized once to reconstruct a mutually exclusive labelled coverage. Every class pair is tested for positive-area intersection, and the stage fails before publication if overlap is detected. The native categorical raster remains authoritative for exact pixel counts. The GeoJSON provenance receipt records source dimensions and hashes, selected pyramid level, level-zero scaling, connectivity, repair diagnostics, requested and selected simplification tolerance, categorical fidelity, segment reduction, selective-smoothing settings, displacement and overlap tests.",
    )

    methods_anchor = find_exact(document, "Data availability")
    insert_before(methods_anchor, "Spatial grid resolution development benchmark", "Heading 2")
    insert_before(
        methods_anchor,
        "For grid-based tissue-domain analysis, the 224 x 224-pixel UNI-2 context field remained physically calibrated to 0.25 micrometres per pixel. Grid-centre stride and the central patch-token pooling square were varied independently. We compared 28/56 and 56/56 inner-square/stride geometries with the historical 90/90 geometry while retaining UNI2-h, GrandQC support, downstream algorithms and KODAMA.matrix ncomp=50. The 28/56 configuration was selected on the three-class development specimen. The supplied expert GeoJSON was rasterized only in the evaluation program after every candidate was complete. Cluster identities were matched to reference identities with a Hungarian assignment. Reported endpoints included class Dice and IoU, tissue-footprint agreement, symmetric internal-boundary distance, boundary F1 at fixed native-pixel tolerances, boundary-length ratio and runtime. Because the reference contained self-intersections and overlapping features, rasterization followed the declared feature order and recorded the overlap count. No reference-derived mask, distance transform, polygon or parameter was passed to UNI-2, KODAMA, clustering, MedSAM or vector generation.",
        "Normal",
    )
    insert_before(methods_anchor, "Selective shared-boundary regularization", "Heading 2")
    insert_before(
        methods_anchor,
        "Selective regularization was applied after fidelity-constrained coverage simplification. Internal shared-boundary vertices were eligible only when their turn was at least 45 degrees and each adjacent edge was at least 16 native pixels. One endpoint-preserving Laplacian pass moved an eligible vertex by 0.05 times its local Laplacian displacement. Vertices on the external tissue boundary were fixed, and the method added no vertices. The complete modified seam network was polygonized once and assigned back to the unique source class, preventing neighbouring classes from being simplified independently. Candidate generation used only the pipeline prediction. The expert annotation was used later to choose among completed development candidates and to quantify the resulting trade-off; the selected constants therefore require confirmation on held-out specimens.",
        "Normal",
    )

    find_exact(document, "Discussion").paragraph_format.page_break_before = True
    normalize_heading_presentation(document)

    set_revision_metadata(
        document,
        subject="Nature Methods Article manuscript",
        description=(
            "Revised 2026-09-15 to include the development selection of UNI2 28/56 spatial-grid "
            "sampling and selective topology-preserving shared-boundary regularization."
        ),
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def revise_supplement(source: Path, destination: Path) -> None:
    document = Document(source)

    replace_starting(
        document,
        "Scope. This document records",
        "Scope. This document records the CellPhenotyper implementation revision through 15 September 2026, its software-level verification, the completed single-specimen development comparison of UNI-2 grid resolution and boundary regularization, and the experiments still required before broader biological or performance claims are made.",
    )
    replace_starting(
        document,
        "Representation and clustering ablation.",
        "Representation and clustering ablation. A single-specimen K=3 development comparison of three spatial-grid geometries is complete and is reported in Supplementary Note 15. Independent specimens are still required to compare cell-centred and spatial-grid observations, token-subset and masked-forward definitions, repeated seeds and clustering hyperparameters. Future reports should include runtime, memory, adjusted Rand index or variation of information across seeds, spatial coherence and reference-standard boundary metrics where available.",
    )
    replace_exact(
        document,
        "Supplementary Note 11 | Prespecified held-out comparison",
        "Supplementary Note 11 | Development reference separation and held-out confirmation",
    )
    replace_starting(
        document,
        "Before revealing the final reference annotation",
        "The initial UNI-2 tiling candidates were generated without access to the expert annotation. After the annotation became available, a separate benchmark program rasterized the completed candidate GeoJSONs, matched arbitrary cluster integers to reference integers and calculated the declared overlap and boundary metrics. The same evaluation-only route was used to select the final shared-boundary regularizer from completed candidates. The production inference programs and Nextflow parameter schema do not accept a reference-annotation path. The selected 28/56 grid geometry and smoothing constants are development-set choices and must be frozen before confirmation on independent held-out specimens. Any future encoder comparison must preserve KODAMA.matrix ncomp=50, clustering policy, GrandQC support, physical scale and evaluation code unless the protocol declares the change in advance.",
    )
    replace_starting(
        document,
        "MASK_TO_GEOJSON converts the final MedSAM-refined",
        "MASK_TO_GEOJSON converts the final MedSAM-refined categorical mask through four-connected rasterio polygonization at native resolution. Invalid point-touching rings are split into valid positive-area parts without changing their class. When shared boundaries contain unequal collinear vertices, combined boundary linework is noded once and polygonized into atomic faces; each face is assigned to the unique covering source class, while background faces are discarded. Every positive-area annotated component is retained. Visvalingam-Whyatt simplification is applied to the complete polygon coverage, so shared edges are simplified once rather than independently for each class. The configured tolerance is an upper bound: binary search selects the strongest candidate whose exact categorical symmetric-difference area does not exceed the declared fraction of native foreground. Interior rings remain present; per-class hole filling and buffer smoothing remain disabled. After simplification, the optional selective pass described in Supplementary Note 16 moves only eligible internal shared-boundary vertices. The vectorizer tests every pairwise positive-area intersection after construction and fails closed if any overlap exceeds numerical tolerance. Its provenance JSON records native source shape and hashes, label-wise pixel counts, selected pyramid level, scaling, connectivity, repair diagnostics, requested and selected tolerance, coordinate and segment reduction, measured fidelity, smoothing parameters and displacement, feature count and mutual exclusivity. The native raster remains the authority for exact categorical counts.",
    )
    replace_starting(
        document,
        "The pre-MedSAM competition step treats",
        "The pre-MedSAM competition step treats the grown KODAMA mask as a categorical partition rather than refining each class independently. An internal boundary is detected where two positive labels meet, dilated to the declared error radius and intersected with the clustering-uncertainty editable mask. Robust Lab and optical-density features are standardized by the tissue median and interquartile range. Label prototypes are estimated from fixed label interiors outside the editable band. The local energy contains squared prototype distance and a Potts disagreement penalty whose coupling decreases across strong image-feature gradients. The default eight-neighbour system includes inverse-distance diagonal interactions to reduce axis-aligned staircasing. Mean-field label probabilities are updated with a Boltzmann temperature that decreases geometrically from 2.0 to 0.05 over 16 iterations. Candidate labels are restricted to the current label and labels on the local frontier. Protected cores are clamped at every iteration, zero-labelled pixels are not created or filled, and all pixels outside GrandQC clean tissue remain zero. The lowest hard-energy state encountered is supplied to MedSAM as its multiclass baseline; the original confident seeds remain the prompt source. The default fourfold working scale bounds WSI memory. This deterministic algorithm uses no expert annotation. Its boundary radius, appearance weight, smoothness weight, edge decay and temperature schedule require prespecified held-out evaluation before accuracy claims.",
    )

    anchor = find_exact(document, "Supplementary reporting checklist")
    sections = [
        (
            "Supplementary Note 15 | UNI2 spatial grid resolution development",
            "The spatial-grid route separates the central token-pooling width from the grid-centre stride while retaining a 224 x 224-pixel UNI2-h context field and 0.25-micrometre-per-pixel model calibration. At source MPP m, a model-space distance d is mapped to round(d x 0.25 / m) level-0 pixels. The development comparison held GrandQC clean-tissue support, UNI2-h weights and preprocessing, KODAMA.matrix ncomp=50, K=3 clustering and downstream reconstruction constant. The tested inner-square/stride pairs were 28/56, 56/56 and the historical coupled 90/90 geometry. The two 56-pixel-stride routes each retained 498,992 observations. At 16-fold reference rasterization, mean class IoU was 0.945360, 0.883707 and 0.681739, respectively; mean class Dice was 0.971665, 0.938255 and 0.8077. Mean symmetric internal-boundary error was 10.81, 18.52 and 37.64 native pixels. Boundary F1 at 32 native pixels was 0.91979, 0.86502 and 0.7329. Summed Nextflow task realtime was 35,247.2 s for 28/56 and 34,925.6 s for 56/56. The 28/56 route was therefore selected for subsequent grid-based tissue-domain development. This is parameter selection on one annotated specimen, not independent validation.",
        ),
        (
            "Supplementary Note 16 | Selective shared boundary regularization",
            "The final vector regularizer operates only after topology-preserving coverage simplification. A shared internal vertex is eligible when its turn angle is at least 45 degrees and both incident segments are at least 16 native pixels. One endpoint-preserving Laplacian pass moves each eligible vertex by coefficient 0.05. Vertices on the external tissue boundary remain fixed, and no vertices are added. Modified seams are polygonized once into atomic faces and reassigned to their source class, after which the pairwise positive-area overlap test must remain zero. On the K=3 development output, fidelity-constrained simplification selected a 14.953125-pixel tolerance, reducing 8,104,295 native raster-edge segments by 96.99% while remaining below the 1% categorical-change budget. Selective regularization retained three features and 244,005 segments, changed no external tissue-footprint area, moved no vertex by more than 11.856 native pixels and preserved zero class overlap. The 95th percentile turn fell from 121.130 to 116.581 degrees, and long sharp-turn density fell from 1,355.00 to 1,292.39 per 100,000 ring pixels. Mean class IoU changed from 0.945360 to 0.942344, and boundary F1 at 32 pixels changed from 0.919790 to 0.919382. The procedure was selected to reduce implausible angular seams with a small measured reference-agreement trade-off; it is not presented as an accuracy gain.",
        ),
        (
            "Supplementary Note 17 | Expert annotation use and limitations",
            "The three-class expert GeoJSON was not stored as a production input and was never passed to UNI-2, KODAMA, clustering, MedSAM, shared-boundary simplification or smoothing. Candidate generation used only the image, GrandQC support and predicted labels. The annotation was read later by a separate benchmark for development ranking and parameter selection. Because the source GeoJSON contained self-intersections and overlapping features, the benchmark rasterized it at a declared 16-fold scale and resolved overlaps by feature order, recording 78,409 overlapping analysis pixels. A Hungarian assignment aligned arbitrary prediction and reference cluster integers before scoring. Metrics included tissue Dice and IoU, class Dice and IoU, categorical accuracy on common tissue, symmetric boundary distance, boundary F1 at 32, 64 and 128 native pixels, boundary-length ratio and angular geometry. The selected constants may overfit this specimen and should be confirmed without adjustment on independent tissue types, staining conditions and scanners.",
        ),
    ]
    for heading, body in sections:
        insert_before(anchor, heading, "Heading 1")
        insert_before(anchor, body, "Normal")

    normalize_heading_presentation(document)

    set_revision_metadata(
        document,
        subject="Nature Methods Supplementary Information",
        description=(
            "Revised 2026-09-15 to document the UNI2 28/56 development comparison, "
            "selective shared-boundary regularization and strict separation of expert evaluation from inference."
        ),
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--main-source", type=Path, required=True)
    parser.add_argument("--supplement-source", type=Path, required=True)
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--date-tag", default="20260915")
    args = parser.parse_args()

    revise_main(
        args.main_source,
        args.outdir / f"CellPhenotyper_Nature_Methods_{args.date_tag}.docx",
    )
    revise_supplement(
        args.supplement_source,
        args.outdir / f"CellPhenotyper_Supplementary_Information_{args.date_tag}.docx",
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
